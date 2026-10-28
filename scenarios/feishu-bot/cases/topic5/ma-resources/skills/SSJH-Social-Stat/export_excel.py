"""
export_excel.py — datashuoshuo-socialmedia Excel 导出 tool
从 run() 返回的 full JSON 或本地 JSON 文件生成 Excel 报表。

用法：
    from export_excel import export_excel
    path = export_excel(result, output_path="/tmp/report.xlsx")
    # 或从文件加载
    path = export_excel("/tmp/datashuoshuo/小米_xxx.json")

version: 1.1.0
"""

import json
import os
import datetime
from typing import Union

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    raise ImportError("需要安装 openpyxl: pip install openpyxl")


# ─────────────────────────────────────────────
# 样式常量
# ─────────────────────────────────────────────
H_FILL  = PatternFill("solid", fgColor="1F4E79")
H_FONT  = Font(color="FFFFFF", bold=True)
SH_FILL = PatternFill("solid", fgColor="D6E4F0")
SH_FONT = Font(bold=True)
THIN    = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)


def _hdr(ws, cols, row=1, fill=H_FILL, font=H_FONT):
    for c, v in enumerate(cols, 1):
        cell = ws.cell(row, c, v)
        cell.fill, cell.font = fill, font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN


def _w(ws, r, c, v, bold=False, wrap=True):
    cell = ws.cell(r, c, v)
    if bold:
        cell.font = Font(bold=True)
    cell.alignment = Alignment(vertical="center", wrap_text=wrap)
    cell.border = THIN
    return cell


def _auto_width(ws, min_w=10, max_w=50):
    for col in ws.columns:
        length = max((len(str(c.value or "")) for c in col), default=min_w)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(length, min_w), max_w)


def _fmt_ts(val):
    if isinstance(val, (int, float)) and val > 1e10:
        return datetime.datetime.fromtimestamp(val / 1000).strftime("%Y-%m-%d %H:%M:%S")
    return val or ""


def _fmt_sentiment(val):
    return {"1": "正面", "0": "中性", "-1": "负面",
            1: "正面", 0: "中性", -1: "负面"}.get(val, str(val) if val is not None else "")


# ─────────────────────────────────────────────
# 各 Sheet 写入函数
# ─────────────────────────────────────────────

def _sheet_overview(wb, result):
    ws = wb.create_sheet("概览")
    _hdr(ws, ["指标", "数值", "说明"])
    meta = result.get("meta", {})
    rows = [
        ("关键词",    meta.get("keyword", ""),              ""),
        ("排除词",    meta.get("filterword", "") or "无",   ""),
        ("平台",      ", ".join(meta.get("platforms", [])), ""),
        ("时间范围",  f"{meta.get('start_date','')} ~ {meta.get('end_date','')}", ""),
        ("内容模式",  meta.get("content_mode", ""),         "UGC/PGC/UKN/ALL"),
        ("帖子类型",  str(meta.get("msg_type", "")),        "0=主+评 1=主 2=评"),
        ("过滤广告",  "是" if meta.get("filter_ad") else "否", ""),
        ("最低互动量", str(meta.get("min_interaction") or "不限"), ""),
        ("原文排序",  meta.get("sort", ""),                 ""),
        ("总声量",    result.get("total_volume"),           ""),
        ("总互动量",  result.get("total_interaction"),      "不支持互动量的平台返回 None"),
        ("净情感度",  result.get("total_sentiment"),        ""),
    ]
    for i, (a, b, c) in enumerate(rows, 2):
        _w(ws, i, 1, a, bold=True); _w(ws, i, 2, b); _w(ws, i, 3, c)
    _auto_width(ws)


def _sheet_data_type(wb, result):
    ws = wb.create_sheet("数据类型分布")
    _hdr(ws, ["平台(source)", "声量数值", "占比%"])
    dt = result.get("data_type", {})
    total = sum(v for v in dt.values() if isinstance(v, (int, float)))
    for i, (k, v) in enumerate(sorted(dt.items(), key=lambda x: x[1], reverse=True), 2):
        pct = round(v / total * 100, 1) if total else 0
        _w(ws, i, 1, k); _w(ws, i, 2, v); _w(ws, i, 3, f"{pct}%")
    _auto_width(ws)


def _sheet_sentiment(wb, result):
    ws = wb.create_sheet("情感分布")
    _hdr(ws, ["平台", "正面", "正面%", "中性", "中性%", "负面", "负面%", "合计"])
    sd = result.get("sentiment_dist", {})
    row = 2
    for plat, vals in sd.items():
        pos = vals.get("positive", 0)
        neu = vals.get("neutral", 0)
        neg = vals.get("negative", 0)
        total = pos + neu + neg
        def p(v): return f"{round(v/total*100,1)}%" if total else "0%"
        label = "汇总" if plat == "_total" else plat
        bold = (plat == "_total")
        _w(ws, row, 1, label, bold=bold)
        _w(ws, row, 2, pos); _w(ws, row, 3, p(pos))
        _w(ws, row, 4, neu); _w(ws, row, 5, p(neu))
        _w(ws, row, 6, neg); _w(ws, row, 7, p(neg))
        _w(ws, row, 8, total)
        row += 1
    _auto_width(ws)


def _sheet_volume_type(wb, result):
    ws = wb.create_sheet("声量类型")
    _hdr(ws, ["平台", "UGC", "UGC%", "PGC", "PGC%", "UKN", "UKN%", "合计"])
    vt = result.get("volume_type", {})
    row = 2
    for plat, vals in vt.items():
        ugc = vals.get("UGC", 0); pgc = vals.get("PGC", 0); ukn = vals.get("UKN", 0)
        total = ugc + pgc + ukn
        def p(v): return f"{round(v/total*100,1)}%" if total else "0%"
        label = "汇总" if plat == "_total" else plat
        bold = (plat == "_total")
        _w(ws, row, 1, label, bold=bold)
        _w(ws, row, 2, ugc); _w(ws, row, 3, p(ugc))
        _w(ws, row, 4, pgc); _w(ws, row, 5, p(pgc))
        _w(ws, row, 6, ukn); _w(ws, row, 7, p(ukn))
        _w(ws, row, 8, total)
        row += 1
    _auto_width(ws)


def _sheet_trend(wb, result):
    """声量趋势 + 互动量趋势合并为一个 sheet：日期 / 平台 / 声量 / 互动量"""
    ws = wb.create_sheet("趋势")
    _hdr(ws, ["日期", "平台", "声量", "互动量"])
    # 收集所有日期+平台组合
    volume_data = result.get("volume_trend") or {}
    interaction_data = result.get("interaction_trend") or {}
    # 合并所有 (date, plat) key
    keys = set()
    for plat, days in volume_data.items():
        for date in (days or {}):
            keys.add((date, plat))
    for plat, days in interaction_data.items():
        for date in (days or {}):
            keys.add((date, plat))
    row = 2
    for date, plat in sorted(keys):
        vol = (volume_data.get(plat) or {}).get(date, "")
        itr = (interaction_data.get(plat) or {}).get(date, "")
        _w(ws, row, 1, date); _w(ws, row, 2, plat)
        _w(ws, row, 3, vol); _w(ws, row, 4, itr)
        row += 1
    _auto_width(ws)


def _sheet_cloud_word(wb, result):
    ws = wb.create_sheet("关键词云")
    _hdr(ws, ["关键词", "平台", "词频"])
    row = 2
    for plat, words in (result.get("cloud_word") or {}).items():
        for word, cnt in sorted(words.items(), key=lambda x: x[1], reverse=True):
            _w(ws, row, 1, word); _w(ws, row, 2, plat); _w(ws, row, 3, cnt)
            row += 1
    _auto_width(ws)


def _sheet_articles(wb, result):
    ws = wb.create_sheet("原文内容")
    fields = [
        "site_name", "publish_time", "title", "content", "url",
        "sentiment", "interaction_cnt", "like_cnt", "review_cnt",
        "reposts_cnt", "fav_cnt", "is_main_post", "is_ad",
        "content_mode", "data_type", "uid", "author",
        "keywords", "note_stickers", "pic_urls", "cover_ocr_content",
    ]
    cn_headers = [
        "平台", "发布时间", "标题", "正文", "链接",
        "情感", "互动量", "点赞数", "评论数",
        "转发数", "收藏数", "主贴/评论", "是否广告",
        "内容类型", "数据类型", "作者UID", "作者名",
        "关键词列表", "话题标签", "图片URL列表", "封面OCR内容",
    ]
    _hdr(ws, cn_headers)
    # 合并所有平台，按互动量全局降序排序
    all_arts = []
    for arts in (result.get("articles") or {}).values():
        all_arts.extend(arts or [])
    all_arts.sort(key=lambda x: x.get("interaction_cnt") or 0, reverse=True)
    row = 2
    for art in all_arts:
        for j, f in enumerate(fields, 1):
            val = art.get(f, "")
            if f == "publish_time":
                val = _fmt_ts(val)
            elif f == "sentiment":
                val = _fmt_sentiment(val)
            elif f in ("keywords", "note_stickers") and isinstance(val, list):
                val = ", ".join(str(v) for v in val)
            elif f == "pic_urls" and isinstance(val, list):
                val = "\n".join(str(v) for v in val)
            _w(ws, row, j, str(val) if val else "", wrap=(f == "content"))
        row += 1
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 50
    ws.column_dimensions["E"].width = 40


def _sheet_region(wb, result):
    ws = wb.create_sheet("区域分布")
    _hdr(ws, ["地区", "声量"])
    rd = result.get("region_dist") or {}
    for i, (k, v) in enumerate(sorted(rd.items(), key=lambda x: x[1], reverse=True), 2):
        _w(ws, i, 1, k); _w(ws, i, 2, v)
    _auto_width(ws)


# ─────────────────────────────────────────────
# 完整性校验（Fix #4：拒绝把摘要形态数据导出为正式 Excel）
# ─────────────────────────────────────────────

_SUMMARY_ONLY_KEYS = ("top_keywords", "article_digest", "top_regions")
_FULL_KEYS = ("cloud_word", "articles", "region_dist")


def _check_full_shape(result: dict) -> None:
    if any(k in result for k in _SUMMARY_ONLY_KEYS) and not any(k in result for k in _FULL_KEYS):
        raise ValueError(
            "检测到摘要数据（mode=summary 或未执行 save），禁止导出为正式 Excel，"
            "请使用完整结果（mode=full）重新执行"
        )


# ─────────────────────────────────────────────
# 主入口
# ─────────────────────────────────────────────

def export_excel(
    data: Union[dict, str],
    output_path: str = None,
) -> str:
    """
    将 full JSON 结果导出为 Excel。

    参数：
        data: run() 返回的 full dict，或本地 JSON 文件路径（str）
        output_path: 输出路径（含文件名），默认与 JSON 同目录

    返回：
        Excel 文件的绝对路径
    """
    # 加载数据
    if isinstance(data, str):
        with open(data, "r", encoding="utf-8") as f:
            result = json.load(f)
        json_path = data
    else:
        result = data
        json_path = result.get("_json_path", "")

    _check_full_shape(result)

    # 输出路径
    if output_path is None:
        if json_path:
            output_path = os.path.splitext(json_path)[0] + ".xlsx"
        else:
            meta = result.get("meta", {})
            kw   = meta.get("keyword", "output").replace("+", "_").replace(" ", "_")
            ts   = datetime.datetime.now().strftime("%Y%m%d%H%M%S")
            output_path = f"/tmp/datashuoshuo/{kw}_{ts}.xlsx"
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    # 构建 Excel
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _sheet_overview(wb, result)
    _sheet_data_type(wb, result)
    _sheet_sentiment(wb, result)
    _sheet_volume_type(wb, result)
    _sheet_trend(wb, result)
    _sheet_cloud_word(wb, result)
    _sheet_articles(wb, result)
    _sheet_region(wb, result)

    wb.save(output_path)
    return output_path
