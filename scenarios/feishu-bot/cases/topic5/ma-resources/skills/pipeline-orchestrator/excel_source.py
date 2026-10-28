"""用户上传活动明细 Excel → clean_campaign 行结构适配。

权威样例列来自 ``瑞幸暹罗厘普联名数据.xlsx``。文件由 Host Gateway 上传 Ark Files API，
并挂载到 ``/mnt/session/uploads/<mount_path>``；confirmed_plan.source_file.mount_path
是唯一运行期引用。没有 source_file 的历史任务返回 None，让 handler 回退 SSJH articles；
已声明文件但不可读则抛错，禁止静默换源。
"""
from __future__ import annotations

import os
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Optional

_CST = timezone(timedelta(hours=8))
_PLATFORM_MAP = {
    "小红书": "xiaohongshu",
    "抖音": "douyin",
    "微博": "weibo",
    "哔哩哔哩": "bilibili",
    "B站": "bilibili",
    "微信": "weixin",
    "微信公众号": "weixin",
}
_CONTENT_MODE = {"UGC": "UGC", "PGC": "PGC", "未知": "UKN", "UKN": "UKN"}


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _integer(value: Any) -> int:
    if value in (None, "", "-", "--"):
        return 0
    try:
        return int(float(str(value).replace(",", "").strip()))
    except (TypeError, ValueError):
        return 0


def _publish_time_ms(value: Any) -> Optional[int]:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        dt = value
    else:
        text = str(value).strip()
        dt = None
        for fmt in (
            "%Y-%m-%d %H:%M:%S",
            "%Y/%m/%d %H:%M:%S",
            "%Y-%m-%d",
            "%Y/%m/%d",
        ):
            try:
                dt = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if dt is None:
            try:
                numeric = int(float(text))
            except (TypeError, ValueError):
                return None
            return numeric if numeric > 10**12 else numeric * 1000
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=_CST)
    return int(dt.timestamp() * 1000)


def _keywords(value: Any) -> list[str]:
    return list(
        dict.fromkeys(
            item.strip().lstrip("#")
            for item in re.split(r"[|,，、\s]+", _text(value))
            if item.strip().lstrip("#")
        )
    )


def normalize_excel_row(row: dict[str, Any]) -> dict:
    """把一行中文 Excel 列转换为 pipeline 统一行结构。"""
    source_content_mode = _text(row.get("内容类型"))
    content_mode = _CONTENT_MODE.get(source_content_mode.upper(), "UKN")
    site_name = _text(row.get("站点名称"))
    interaction = _integer(row.get("总互动量"))
    if not interaction:
        interaction = sum(
            _integer(row.get(key))
            for key in ("转发数", "评论数", "点赞数", "收藏数")
        )
    item_id = _text(row.get("id"))
    return {
        "id": item_id,
        "fingerprint": item_id,
        "data_type": _text(row.get("数据类型")),
        "analysis_object": _text(row.get("分析对象")),
        "site_name": site_name,
        "_platform": _PLATFORM_MAP.get(site_name, site_name or "unknown"),
        "title": _text(row.get("标题")),
        "content": _text(row.get("内容")),
        "video_cover_text": _text(row.get("视频封面内容识别")),
        "audio_text": _text(row.get("音频内容识别")),
        "video_flower_text": _text(row.get("视频花字识别")),
        "video_text": _text(row.get("视频文本识别")),
        "content_mode": content_mode,
        "_source_content_mode": source_content_mode,
        "keywords": _keywords(row.get("话题")),
        "author": _text(row.get("作者")),
        "url": _text(row.get("url")),
        "publish_time": _publish_time_ms(row.get("发表时间")),
        "interaction_cnt": interaction,
        "share_count": _integer(row.get("转发数")),
        "comment_count": _integer(row.get("评论数")),
        "like_count": _integer(row.get("点赞数")),
        "read_count": _integer(row.get("阅读数")),
        "collect_count": _integer(row.get("收藏数")),
        "fans": _integer(row.get("粉丝数")),
        "_source": "user_uploaded_excel",
    }


def resolve_activity_xlsx(plan: dict) -> Optional[Path]:
    """解析已冻结的 Session 挂载路径；测试可用 TOPIC5_ACTIVITY_XLSX_PATH 覆盖。"""
    override = os.environ.get("TOPIC5_ACTIVITY_XLSX_PATH")
    if override:
        return Path(override).expanduser()
    source = plan.get("source_file") if isinstance(plan, dict) else None
    if not isinstance(source, dict) or not source.get("mount_path"):
        return None
    mount_path = str(source["mount_path"])
    if Path(mount_path).name != mount_path or not mount_path.lower().endswith(".xlsx"):
        raise ValueError("excel_source: source_file.mount_path 非法")
    uploads_dir = Path(os.environ.get("TOPIC5_SESSION_UPLOADS_DIR", "/mnt/session/uploads"))
    return uploads_dir / mount_path


def load_activity_rows(plan: dict) -> Optional[dict]:
    """读取活动明细；无 source_file 返回 None，声明后缺文件/坏文件则报错。"""
    path = resolve_activity_xlsx(plan)
    if path is None:
        return None
    if not path.is_file():
        raise FileNotFoundError(f"excel_source: 已挂载活动明细不存在: {path}")

    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("excel_source: 运行环境缺少 openpyxl") from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header_values = next(rows, None)
        if not header_values:
            raise ValueError("excel_source: Excel 没有表头")
        headers = [_text(value) for value in header_values]
        required = {
            "站点名称", "标题", "内容", "发表时间", "内容类型",
            "分析对象", "url",
        }
        missing = sorted(required - set(headers))
        if missing:
            raise ValueError("excel_source: 缺少必要列: " + "、".join(missing))
        interaction_columns = {
            "总互动量", "转发数", "评论数", "点赞数", "收藏数",
        }
        if not interaction_columns.intersection(headers):
            raise ValueError("excel_source: 缺少互动量列（总互动量或互动分项）")

        content_scope = str(plan.get("content_scope") or "all").lower()
        expected_mode = {"ugc": "UGC", "pgc": "PGC"}.get(content_scope)
        expected_platforms = {
            _PLATFORM_MAP.get(str(value), str(value))
            for value in (plan.get("platforms") or [])
        }
        expected_campaigns = {
            str(item.get("id"))
            for item in (plan.get("campaigns") or [])
            if isinstance(item, dict) and item.get("id")
        }
        period_start = str(plan.get("period_start") or "").replace("-", "")
        period_end = str(plan.get("period_end") or "").replace("-", "")
        normalized: list[dict] = []
        source_row_count = 0
        for row_number, values in enumerate(rows, 2):
            source_row_count += 1
            raw = dict(zip(headers, values))
            row = normalize_excel_row(raw)
            if not any((row["id"], row["title"], row["content"], row["url"])):
                continue
            missing_values = [
                key for key, value in (
                    ("站点名称", row["_platform"]),
                    ("标题/正文", row["title"] or row["content"]),
                    ("发表时间", row["publish_time"]),
                    ("内容类型", row["_source_content_mode"]),
                    ("分析对象", row["analysis_object"]),
                    ("URL", row["url"]),
                )
                if not value or value == "unknown"
            ]
            if missing_values:
                raise ValueError(
                    f"excel_source: 第 {row_number} 行缺少或无法解析: "
                    + "、".join(missing_values)
                )
            if expected_platforms and row["_platform"] not in expected_platforms:
                raise ValueError(
                    f"excel_source: 第 {row_number} 行平台 {row['_platform']} "
                    "不在确认范围"
                )
            publish_ymd = datetime.fromtimestamp(
                row["publish_time"] / 1000, tz=_CST
            ).strftime("%Y%m%d")
            if (
                period_start and period_end
                and not (period_start <= publish_ymd <= period_end)
            ):
                raise ValueError(
                    f"excel_source: 第 {row_number} 行发表时间 {publish_ymd} "
                    "超出确认周期"
                )
            if expected_mode and row["content_mode"] != expected_mode:
                continue
            normalized.append(row)
    finally:
        workbook.close()

    if not normalized:
        raise ValueError("excel_source: 按内容范围过滤后没有可用活动明细")
    invalid_content_modes: dict[str, int] = {}
    for row in normalized:
        source_mode = str(row.get("_source_content_mode") or "").strip()
        if source_mode.upper() not in _CONTENT_MODE:
            invalid_content_modes[source_mode] = (
                invalid_content_modes.get(source_mode, 0) + 1
            )
    invalid_content_mode_count = sum(invalid_content_modes.values())
    if (
        invalid_content_mode_count
        and invalid_content_mode_count / source_row_count >= 0.05
    ):
        raise ValueError(
            "excel_source: 内容类型出现大量非 PGC/UGC/未知取值，"
            f"占比={invalid_content_mode_count / source_row_count:.2%}，"
            f"明细={invalid_content_modes}"
        )
    source_campaigns = {
        str(row.get("analysis_object") or "").strip()
        for row in normalized
        if str(row.get("analysis_object") or "").strip()
    }
    campaign_mapping = {value: value for value in source_campaigns}
    if expected_campaigns and source_campaigns != expected_campaigns:
        if len(expected_campaigns) == 1 and len(source_campaigns) == 1:
            source_id = next(iter(source_campaigns))
            confirmed_id = next(iter(expected_campaigns))
            campaign_mapping = {source_id: confirmed_id}
            for row in normalized:
                row["_source_analysis_object"] = row["analysis_object"]
                row["analysis_object"] = confirmed_id
        else:
            raise ValueError(
                "excel_source: Excel 分析对象与确认活动分组不一致；"
                f"Excel={sorted(source_campaigns)}，确认={sorted(expected_campaigns)}"
            )
    return {
        "rows": normalized,
        "meta": {
            "source": "user_uploaded_excel",
            "file_name": (plan.get("source_file") or {}).get("name") or path.name,
            "path": str(path),
            "content_scope": str(plan.get("content_scope") or "all").lower(),
            "platforms": sorted({row["_platform"] for row in normalized}),
            "source_row_count": source_row_count,
            "accepted_row_count": len(normalized),
            "invalid_content_modes_mapped_to_unknown": invalid_content_modes,
            "campaigns": sorted({row["analysis_object"] for row in normalized}),
            "source_campaigns": sorted(source_campaigns),
            "campaign_mapping": campaign_mapping,
            "period_start": period_start,
            "period_end": period_end,
        },
    }
